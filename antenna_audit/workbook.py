"""Render a :class:`SheetPlan` into an .xlsx workbook."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import layout
from .imaging import ImagePreparer
from .plan import SheetPlan, SlotPlan

# Palette, kept close to the reference workbook's restrained grey styling.
INK = "FF1F2933"
MUTED = "FF4D5761"
FAINT = "FF9AA5B1"
BANNER_FILL = "FF1F2933"
BANNER_TEXT = "FFFFFFFF"
ZONE_FILL = "FFE4E7EB"
PRE_FILL = "FFDCEBDC"
BEFORE_FILL = "FFE7E3F2"
POST_FILL = "FFF6E7DC"
PLACEHOLDER_FILL = "FFFAFBFC"
MISSING_FILL = "FFFDF3F3"

TITLE_FONT = Font(name="Calibri", size=16, bold=True, color=INK)
SUBTITLE_FONT = Font(name="Calibri", size=9, color=MUTED)
BANNER_FONT = Font(name="Calibri", size=12, bold=True, color=BANNER_TEXT)
ZONE_FONT = Font(name="Calibri", size=10, bold=True, color=INK)
PREPOST_FONT = Font(name="Calibri", size=10, bold=True, color=INK)
HEADING_FONT = Font(name="Calibri", size=10, bold=True, color=MUTED)
PLACEHOLDER_FONT = Font(name="Calibri", size=9, italic=True, color=FAINT)
MISSING_FONT = Font(name="Calibri", size=9, italic=True, color="FFB04A4A")
VALUES_HEADER_FONT = Font(name="Calibri", size=11, bold=True, color=INK)

# The Values sheet, which the field team fills in beside the photos.  Its row
# labels are worded differently from the photo headings ("850 E Tilt", not
# "850 Tilt"; "Antenna Azimuth", not "Antenna Azimuth Photo"), so they are kept
# separate rather than derived from catalog.DISPLAY_NAMES.
VALUES_SHEET_TITLE = "Values"
VALUES_COLUMNS = ("sector", None, "Pre", "Before Swap", "Post", "Plan")
VALUES_ROWS = [
    "Sec {sector}_ 850 E Tilt",
    "Sec {sector}_ 900 E Tilt",
    "Sec {sector}_ 1800 E Tilt 1",
    "Sec {sector}_ 1800 E Tilt 2",
    "Sec {sector}_ 2100 E Tilt",
    "Sec {sector} Antenna M Tilt",
    "Sec {sector} Antenna Azimuth",
]

CENTRE = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")

_DASHED = Side(style="dashed", color=FAINT)
_THIN = Side(style="thin", color=FAINT)
PLACEHOLDER_BORDER = Border(left=_DASHED, right=_DASHED, top=_DASHED, bottom=_DASHED)
PHOTO_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


@dataclass
class BuildResult:
    """What one site's build produced."""

    site: str
    path: Path
    photos_placed: int
    photos_unplaced: dict[str, int]
    missing_slots: list[str]
    failures: list[tuple[Path, str]]
    rows: int


def _ref(col0: int, row: int, cols: int = 1, rows: int = 1) -> str:
    """Build an A1-style range from a zero-based column and one-based row."""
    start = f"{get_column_letter(col0 + 1)}{row}"
    end = f"{get_column_letter(col0 + cols)}{row + rows - 1}"
    return start if start == end else f"{start}:{end}"


def _write_merged(
    ws, col0: int, row: int, cols: int, text: str, font: Font,
    fill: str | None = None, alignment: Alignment = CENTRE,
    border: Border | None = None, rows: int = 1,
) -> None:
    """Write ``text`` into a merged block, styling every cell it covers."""
    ref = _ref(col0, row, cols, rows)
    if cols > 1 or rows > 1:
        ws.merge_cells(ref)
    cell = ws.cell(row=row, column=col0 + 1)
    cell.value = text
    cell.font = font
    cell.alignment = alignment
    if fill:
        pattern = PatternFill("solid", fgColor=fill)
        for r in range(row, row + rows):
            for c in range(col0 + 1, col0 + cols + 1):
                ws.cell(row=r, column=c).fill = pattern
    if border:
        _outline(ws, col0, row, cols, rows, border)


def _outline(ws, col0: int, row: int, cols: int, rows: int, border: Border) -> None:
    """Draw ``border`` around the outside edge of a cell block."""
    for r in range(row, row + rows):
        for c in range(col0 + 1, col0 + cols + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = Border(
                left=border.left if c == col0 + 1 else None,
                right=border.right if c == col0 + cols else None,
                top=border.top if r == row else None,
                bottom=border.bottom if r == row + rows - 1 else None,
            )


def _anchor(box: layout.Box, width: int, height: int) -> OneCellAnchor:
    """Anchor an image of the given pixel size centred inside ``box``.

    A one-cell anchor with an explicit extent pins the top-left corner and fixes
    the drawn size, so the picture keeps its aspect ratio even if a reader
    resizes rows or columns.
    """
    dx, dy = box.centre_offsets(width, height)
    marker = AnchorMarker(
        col=box.col0,
        colOff=dx * layout.EMU_PER_PX,
        row=box.row0 - 1,
        rowOff=dy * layout.EMU_PER_PX,
    )
    size = XDRPositiveSize2D(
        cx=width * layout.EMU_PER_PX, cy=height * layout.EMU_PER_PX
    )
    return OneCellAnchor(_from=marker, ext=size)


def _apply_dimensions(ws, plan: SheetPlan) -> None:
    """Fix every column width and the default row height."""
    ws.sheet_format.defaultRowHeight = layout.px_to_row_height(layout.ROW_PX)
    ws.sheet_format.customHeight = True
    for col0, px in layout.column_widths().items():
        ws.column_dimensions[get_column_letter(col0 + 1)].width = (
            layout.px_to_col_width(px)
        )
    for sector in plan.sectors:
        ws.row_dimensions[sector.banner_row].height = layout.BANNER_ROW_PT
        ws.row_dimensions[sector.zone_header_row].height = layout.ZONE_ROW_PT


def _write_slot(ws, slot: SlotPlan, preparer: ImagePreparer) -> bool:
    """Place one photo, or draw one placeholder box.  Returns True if a photo landed."""
    box = slot.box
    if slot.photo is None:
        missing = slot.kind == "pre"
        _write_merged(
            ws, box.col0, box.row0, box.cols,
            slot.placeholder_text,
            MISSING_FONT if missing else PLACEHOLDER_FONT,
            fill=MISSING_FILL if missing else PLACEHOLDER_FILL,
            border=PLACEHOLDER_BORDER,
            rows=box.rows,
        )
        return False

    prepared = preparer.prepare(slot.photo.path)
    if prepared is None:
        _write_merged(
            ws, box.col0, box.row0, box.cols,
            f"Could not read\n{slot.photo.name}",
            MISSING_FONT, fill=MISSING_FILL,
            border=PLACEHOLDER_BORDER, rows=box.rows,
        )
        return False

    width, height = layout.fit_within(
        prepared.width, prepared.height, box.width_px, box.height_px
    )
    image = XLImage(str(prepared.path))
    image.width, image.height = width, height
    image.anchor = _anchor(box, width, height)
    ws.add_image(image)
    return True


def _write_values_sheet(wb, plan: SheetPlan) -> None:
    """Add the Values sheet: one row per sector per measured quantity.

    Only the scaffold is written.  The readings themselves come off instruments
    in the field, not out of the photographs, so every value cell is left empty
    for the engineer to fill in.
    """
    ws = wb.create_sheet(VALUES_SHEET_TITLE)
    for index, heading in enumerate(VALUES_COLUMNS, start=1):
        if heading is None:
            continue                      # the label column carries no header
        cell = ws.cell(row=1, column=index)
        cell.value = heading
        cell.font = VALUES_HEADER_FONT

    row = 2
    for sector in plan.sectors:
        for template in VALUES_ROWS:
            ws.cell(row=row, column=1).value = f"sec{sector.number}"
            ws.cell(row=row, column=2).value = template.format(
                sector=sector.number
            )
            row += 1

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 25
    for letter in ("C", "D", "E", "F"):
        ws.column_dimensions[letter].width = 14
    ws.column_dimensions["G"].width = 62


def write_workbook(
    plan: SheetPlan, out_path: Path, preparer: ImagePreparer,
    unplaced: dict[str, int] | None = None,
) -> BuildResult:
    """Write ``plan`` to ``out_path`` and return a summary of what happened."""
    wb = Workbook()
    ws = wb.active
    ws.title = plan.site[:31] or "Sheet1"
    ws.sheet_view.zoomScale = 70
    _apply_dimensions(ws, plan)

    span = layout.LAST_COL - layout.MARGIN_LEFT + 1
    _write_merged(
        ws, layout.MARGIN_LEFT, plan.title_row, span,
        f"{plan.site} — Antenna Audit Photos", TITLE_FONT, alignment=LEFT,
    )
    _write_merged(
        ws, layout.MARGIN_LEFT, plan.subtitle_row, span,
        "Pre photos placed automatically from the field survey. "
        "Before Swap and Post photos are pasted into the dashed boxes by hand.",
        SUBTITLE_FONT, alignment=LEFT,
    )

    placed = 0
    for sector in plan.sectors:
        _write_merged(
            ws, layout.MARGIN_LEFT, sector.banner_row, span,
            f"Sector {sector.number}", BANNER_FONT, fill=BANNER_FILL,
        )
        for zone in (sector.left, sector.right):
            zone_cols = (
                zone.post_col0 + layout.BOX_COLS - zone.pre_col0
            )
            _write_merged(
                ws, zone.pre_col0, sector.zone_header_row, zone_cols,
                zone.title, ZONE_FONT, fill=ZONE_FILL,
            )
            for kind, label, fill in (
                ("pre", "Pre", PRE_FILL),
                ("before", "Before Swap", BEFORE_FILL),
                ("post", "Post", POST_FILL),
            ):
                _write_merged(
                    ws, zone.col0_for(kind), sector.prepost_header_row,
                    layout.BOX_COLS, label, PREPOST_FONT, fill=fill,
                )

            for category in zone.categories:
                for col0 in category.heading_cols:
                    _write_merged(
                        ws, col0, category.heading_row, layout.BOX_COLS,
                        category.heading, HEADING_FONT, alignment=LEFT,
                    )
                for slot in category.slots:
                    if _write_slot(ws, slot, preparer):
                        placed += 1

    _write_values_sheet(wb, plan)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)

    return BuildResult(
        site=plan.site,
        path=out_path,
        photos_placed=placed,
        photos_unplaced=unplaced or {},
        missing_slots=plan.missing_slots,
        failures=list(preparer.failures),
        rows=plan.total_rows,
    )
